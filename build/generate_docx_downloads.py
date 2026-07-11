"""Generate accessible Word downloads for the public document pages.

Run with the bundled Python runtime that includes python-docx:

  PYTHONPATH=/path/to/bundled/python /path/to/bundled/python/bin/python3 build/generate_docx_downloads.py
"""
from __future__ import annotations

import html
import pathlib
import re
from typing import Iterable

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches, Pt, RGBColor

import render_site


OUT_DIR = render_site.SITE / "downloads"
FONT = "Arial"


def strip_inline_markdown(text: str) -> str:
    text = re.sub(r"`([^`]+)`", r"\1", text)
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
    text = re.sub(r"\*\*([^*]+)\*\*", r"\1", text)
    text = re.sub(r"(?<!\*)\*([^*\n]+)\*(?!\*)", r"\1", text)
    return html.unescape(text)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    for existing in tc_pr.findall(qn("w:shd")):
        tc_pr.remove(existing)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    tc_pr.insert_element_before(
        shd,
        "w:noWrap",
        "w:tcMar",
        "w:textDirection",
        "w:tcFitText",
        "w:vAlign",
        "w:hideMark",
        "w:headers",
    )


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_hyperlink(paragraph, label: str, url: str):
    part = paragraph.part
    r_id = part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)

    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "4F46E5")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(color)
    r_pr.append(underline)
    run.append(r_pr)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


INLINE_PATTERN = re.compile(
    r"(\[([^\]]+)\]\((https?://[^)]+)\)|`([^`]+)`|\*\*([^*]+)\*\*|(?<!\*)\*([^*\n]+)\*(?!\*))"
)


def add_inline(paragraph, text: str) -> None:
    pos = 0
    for match in INLINE_PATTERN.finditer(text):
        if match.start() > pos:
            paragraph.add_run(html.unescape(text[pos : match.start()]))
        token = match.group(0)
        if match.group(2) and match.group(3):
            add_hyperlink(paragraph, html.unescape(match.group(2)), match.group(3))
        elif match.group(4):
            run = paragraph.add_run(html.unescape(match.group(4)))
            run.font.name = FONT
        elif match.group(5):
            run = paragraph.add_run(html.unescape(match.group(5)))
            run.bold = True
        elif match.group(6):
            run = paragraph.add_run(html.unescape(match.group(6)))
            run.italic = True
        else:
            paragraph.add_run(html.unescape(token))
        pos = match.end()
    if pos < len(text):
        paragraph.add_run(html.unescape(text[pos:]))


def add_paragraph(doc: Document, text: str, style: str | None = None):
    paragraph = doc.add_paragraph(style=style)
    add_inline(paragraph, text)
    return paragraph


def add_table(doc: Document, lines: list[str]) -> None:
    rows = [[cell.strip() for cell in line.strip().strip("|").split("|")] for line in lines]
    if len(rows) < 2:
        return
    headers = rows[0]
    body = rows[2:]
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_repeat_table_header(table.rows[0])
    for index, cell in enumerate(table.rows[0].cells):
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        set_cell_shading(cell, "E7E7FB")
        paragraph = cell.paragraphs[0]
        paragraph.style = doc.styles["Normal"]
        run = paragraph.add_run(strip_inline_markdown(headers[index]))
        run.bold = True
    for row in body:
        row = row + [""] * (len(headers) - len(row))
        cells = table.add_row().cells
        for index, cell in enumerate(cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            add_inline(cell.paragraphs[0], row[index])
    doc.add_paragraph()


def iter_markdown_blocks(markdown: str, title: str) -> Iterable[tuple[str, object]]:
    lines = markdown.splitlines()
    if lines and lines[0].lstrip().startswith("# "):
        first = lines[0].lstrip()[2:].strip()
        if first == title:
            lines = lines[1:]

    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()
        if not stripped:
            i += 1
            continue
        if stripped.startswith("|") and i + 1 < len(lines) and re.match(
            r"^\s*\|?[\s:|-]+\|[\s:|-]+\|?\s*$", lines[i + 1]
        ):
            table_lines = [line]
            i += 1
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i].rstrip())
                i += 1
            yield "table", table_lines
            continue
        heading = re.match(r"^(#{1,6})\s+(.+)$", stripped)
        if heading:
            level = max(1, min(len(heading.group(1)) - 1, 4))
            yield "heading", (level, strip_inline_markdown(heading.group(2).strip()))
            i += 1
            continue
        if stripped.startswith(">"):
            quote_lines = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                quote_lines.append(lines[i].strip().lstrip(">").strip())
                i += 1
            yield "quote", " ".join(quote_lines)
            continue
        if re.match(r"^[-*]\s+", stripped):
            items = []
            while i < len(lines) and re.match(r"^\s*[-*]\s+", lines[i].strip()):
                items.append(re.sub(r"^[-*]\s+", "", lines[i].strip()))
                i += 1
            yield "bullets", items
            continue
        if re.match(r"^\d+\.\s+", stripped):
            items = []
            while i < len(lines) and re.match(r"^\s*\d+\.\s+", lines[i].strip()):
                items.append(re.sub(r"^\d+\.\s+", "", lines[i].strip()))
                i += 1
            yield "numbers", items
            continue
        para = [stripped]
        i += 1
        while i < len(lines):
            nxt = lines[i].strip()
            if (
                not nxt
                or nxt.startswith("#")
                or nxt.startswith(">")
                or nxt.startswith("|")
                or re.match(r"^[-*]\s+", nxt)
                or re.match(r"^\d+\.\s+", nxt)
            ):
                break
            para.append(nxt)
            i += 1
        yield "paragraph", " ".join(para)


def configure_styles(doc: Document) -> None:
    zoom = doc.settings.element.find(qn("w:zoom"))
    if zoom is not None and zoom.get(qn("w:percent")) is None:
        zoom.set(qn("w:percent"), "100")

    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(8)

    for name, size in [("Title", 22), ("Heading 1", 16), ("Heading 2", 14), ("Heading 3", 12)]:
        style = doc.styles[name]
        style.font.name = FONT
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(21, 22, 26)
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)


def accessibility_note(blocks: list[tuple[str, object]], markdown: str) -> str:
    features = ["built-in heading styles"]
    if any(kind in ("bullets", "numbers") for kind, _ in blocks):
        features.append("list styles")
    if re.search(r"\[[^\]]+\]\(https?://[^)]+\)", markdown):
        features.append("descriptive links")
    if any(kind == "table" for kind, _ in blocks):
        features.append("repeated table header rows")
    if len(features) > 2:
        listed = ", ".join(features[:-1]) + ", and " + features[-1]
    elif len(features) == 2:
        listed = f"{features[0]} and {features[1]}"
    else:
        listed = features[0]
    return f"Accessibility note: this Word document uses {listed} for assistive technology."


def build_docx(doc_meta: dict) -> pathlib.Path:
    markdown = pathlib.Path(doc_meta["source"]).read_text(encoding="utf-8")
    doc = Document()
    configure_styles(doc)
    props = doc.core_properties
    props.title = doc_meta["title"]
    props.subject = doc_meta["source_note"]
    props.author = "AI Infrastructure and Architecture"
    props.keywords = "AI infrastructure, curriculum, accessible document"
    props.comments = "Generated from site markdown with semantic headings, lists, and table headers."

    blocks = list(iter_markdown_blocks(markdown, doc_meta["title"]))

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.add_run(doc_meta["title"])
    add_paragraph(doc, doc_meta["source_note"])
    note = add_paragraph(doc, accessibility_note(blocks, markdown))
    note.runs[0].italic = True

    for kind, payload in blocks:
        if kind == "heading":
            level, text = payload
            doc.add_heading(text, level=level)
        elif kind == "table":
            add_table(doc, payload)
        elif kind == "quote":
            paragraph = add_paragraph(doc, payload)
            paragraph.paragraph_format.left_indent = Inches(0.25)
            for run in paragraph.runs:
                run.italic = True
        elif kind == "bullets":
            for item in payload:
                add_paragraph(doc, item, style="List Bullet")
        elif kind == "numbers":
            for item in payload:
                add_paragraph(doc, item, style="List Number")
        elif kind == "paragraph":
            add_paragraph(doc, payload)

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    out_path = render_site.SITE / doc_meta.get("download", f"downloads/{doc_meta['slug']}.docx")
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)
    return out_path


def main() -> None:
    outputs = [build_docx(doc) for doc in render_site.DOCUMENTS]
    for path in outputs:
        print(path.relative_to(render_site.ROOT))


if __name__ == "__main__":
    main()
